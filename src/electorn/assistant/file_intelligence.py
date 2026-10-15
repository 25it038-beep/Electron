"""Secure file download, inspection, multi-format parsing, and document intelligence."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import mimetypes
import os
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple

from electorn.assistant.schema import DocumentAnalysisReport

logger = logging.getLogger(__name__)


class FileSecurityError(ValueError):
    """Raised on path traversal or unsafe file access."""
    pass


class FileIntelligenceEngine:
    """Safely downloads attachments and performs deep semantic extraction on documents."""

    def __init__(self, download_dir: Optional[Path] = None, max_size_mb: int = 25) -> None:
        self.download_dir = download_dir or (Path.home() / ".electorn" / "assistant" / "downloads")
        self.download_dir.mkdir(parents=True, exist_ok=True)
        self.max_size_bytes = max_size_mb * 1024 * 1024

    def sanitize_path(self, filename: str) -> Path:
        """Prevent directory traversal attacks by normalizing and verifying root boundary."""
        clean_name = re.sub(r'[\\/*?:"<>|]', "_", os.path.basename(filename)).strip()
        if not clean_name:
            clean_name = "unnamed_file.bin"
        target_path = (self.download_dir / clean_name).resolve()
        # Verify target is inside download_dir
        if not str(target_path).startswith(str(self.download_dir.resolve())):
            raise FileSecurityError("Directory traversal detected in filename.")
        return target_path

    def save_attachment_bytes(self, filename: str, content: bytes, source_id: Optional[str] = None) -> Tuple[str, str]:
        """Save file bytes with collision-resistant naming and hash tracking."""
        if len(content) > self.max_size_bytes:
            raise ValueError(f"File size exceeds maximum threshold of {self.max_size_bytes // (1024*1024)} MB.")

        sha256 = hashlib.sha256(content).hexdigest()
        sanitized = self.sanitize_path(filename)

        # Collision avoidance: if exists with different hash, prefix hash
        if sanitized.exists():
            existing_hash = hashlib.sha256(sanitized.read_bytes()).hexdigest()
            if existing_hash != sha256:
                sanitized = self.download_dir / f"{sha256[:8]}_{sanitized.name}"

        sanitized.write_bytes(content)
        return str(sanitized), sha256

    def detect_real_mime(self, file_path: Path) -> str:
        """Inspect file header magic bytes rather than trusting the extension blindly."""
        if not file_path.exists():
            return "application/octet-stream"

        header = file_path.read_bytes()[:16]

        if header.startswith(b"%PDF"):
            return "application/pdf"
        if header.startswith(b"PK\x03\x04"):
            # ZIP container: could be docx, xlsx, or zip
            ext = file_path.suffix.lower()
            if ext in (".docx", ".doc"):
                return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if ext in (".xlsx", ".xls"):
                return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return "application/zip"
        if header.startswith(b"\x89PNG\r\n\x1a\n"):
            return "image/png"
        if header.startswith(b"\xff\xd8\xff"):
            return "image/jpeg"
        if header.startswith(b"GIF87a") or header.startswith(b"GIF89a"):
            return "image/gif"

        ext = file_path.suffix.lower()
        if ext == ".csv":
            return "text/csv"

        # Guess from extension as fallback
        guessed, _ = mimetypes.guess_type(str(file_path))
        return guessed or "text/plain"

    def analyze_file(self, target_file: Path) -> DocumentAnalysisReport:
        """Deep analysis across PDF, DOCX, XLSX, CSV, JSON, TXT, Source Code, and Images."""
        if not target_file.exists():
            raise FileNotFoundError(f"File not found: {target_file}")

        file_size = target_file.stat().st_size
        mime = self.detect_real_mime(target_file)
        ext = target_file.suffix.lower()

        text_content = ""
        key_findings: List[str] = []
        deadlines: List[str] = []
        extracted_data: Dict[str, Any] = {}
        page_count: Optional[int] = None
        limitations: Optional[str] = None
        is_partial = False

        try:
            # 1. PDF EXTRACTION
            if "pdf" in mime or ext == ".pdf":
                try:
                    import pypdf
                    reader = pypdf.PdfReader(str(target_file))
                    page_count = len(reader.pages)
                    pages_text = []
                    for idx, p in enumerate(reader.pages[:50]):  # Protect against memory exhaustion on 1000-page docs
                        t = p.extract_text() or ""
                        if t.strip():
                            pages_text.append(f"[Page {idx+1}]\n{t}")
                    if len(reader.pages) > 50:
                        is_partial = True
                        limitations = f"Document has {len(reader.pages)} pages; first 50 pages analyzed."
                    text_content = "\n\n".join(pages_text)
                    key_findings.append(f"PDF document containing {page_count} page(s).")
                except Exception as e:
                    limitations = f"PDF text extraction limited: {e}"

            # 2. DOCX EXTRACTION
            elif "wordprocessingml" in mime or ext in (".docx", ".doc"):
                try:
                    import docx
                    doc = docx.Document(str(target_file))
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    text_content = "\n".join(paragraphs)
                    tables_count = len(doc.tables)
                    key_findings.append(f"Word document with {len(paragraphs)} paragraph(s) and {tables_count} table(s).")
                except Exception as e:
                    limitations = f"DOCX parsing notice: {e}"

            # 3. SPREADSHEETS (CSV, XLSX)
            elif ext == ".csv":
                try:
                    with open(target_file, "r", encoding="utf-8", errors="replace") as f:
                        reader = csv.reader(f)
                        rows = [r for r in reader]
                        headers = rows[0] if rows else []
                        extracted_data["headers"] = headers
                        extracted_data["total_rows"] = len(rows)
                        key_findings.append(f"CSV dataset containing {len(rows)} rows and columns: {', '.join(headers[:8])}.")
                        text_content = "\n".join([", ".join(r[:10]) for r in rows[:100]])
                except Exception as e:
                    limitations = f"CSV parsing notice: {e}"

            elif "spreadsheetml" in mime or ext in (".xlsx", ".xls"):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(str(target_file), data_only=True, read_only=True)
                    sheet_names = wb.sheetnames
                    extracted_data["sheets"] = sheet_names
                    key_findings.append(f"Excel workbook with sheets: {', '.join(sheet_names)}.")
                    sample_lines = []
                    for s in sheet_names[:3]:
                        ws = wb[s]
                        sample_lines.append(f"--- Sheet: {s} ---")
                        for r in ws.iter_rows(max_row=30, values_only=True):
                            row_vals = [str(c) for c in r if c is not None]
                            if row_vals:
                                sample_lines.append(" | ".join(row_vals[:10]))
                    text_content = "\n".join(sample_lines)
                except Exception as e:
                    limitations = f"Spreadsheet parsing notice: {e}"

            # 4. JSON / CODE / TEXT
            elif ext in (".json", ".txt", ".md", ".py", ".js", ".ts", ".html", ".css", ".yaml", ".yml"):
                raw_text = target_file.read_text(encoding="utf-8", errors="replace")
                text_content = raw_text[:50000] # Safe 50k character cap
                if len(raw_text) > 50000:
                    is_partial = True
                    limitations = "File exceeds 50,000 characters; first 50k analyzed."
                if ext == ".json":
                    try:
                        parsed = json.loads(raw_text)
                        extracted_data["top_keys"] = list(parsed.keys()) if isinstance(parsed, dict) else len(parsed)
                    except Exception:
                        pass
                key_findings.append(f"Text/Source file formatted as {ext.upper()} ({len(raw_text)} characters).")

            # 5. IMAGES
            elif "image" in mime or ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp"):
                try:
                    from PIL import Image
                    with Image.open(str(target_file)) as img:
                        w, h = img.size
                        key_findings.append(f"Image ({img.format}) with dimensions {w}x{h} px in {img.mode} mode.")
                        extracted_data["dimensions"] = f"{w}x{h}"
                        extracted_data["format"] = img.format
                except Exception as e:
                    limitations = f"Image metadata notice: {e}"

        except Exception as err:
            limitations = f"Extraction encountered error: {err}"

        # Extract deadlines, dates, and action items using regex heuristics
        if text_content:
            date_matches = re.findall(
                r"\b(?:due|deadline|by|before|date|scheduled for)\s*[:\-]?\s*([A-Za-z]+\s+\d{1,2}(?:st|nd|rd|th)?|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})",
                text_content,
                re.IGNORECASE
            )
            for dm in set(date_matches[:5]):
                deadlines.append(f"Mentioned date/deadline: {dm}")

            action_matches = re.findall(
                r"\b(?:action item|todo|must|require[ds]? to|please ensure)\s*[:\-]?\s*([^\.\n]{10,120})",
                text_content,
                re.IGNORECASE
            )
            for am in set(action_matches[:5]):
                deadlines.append(f"Action item: {am.strip()}")

        # Synthesize Executive Summary
        summary = (
            f"Analyzed {target_file.name} ({mime}, {round(file_size/1024, 1)} KB). "
            f"Extracted {len(key_findings)} core structural insight(s) and {len(deadlines)} action/deadline indicator(s)."
        )

        return DocumentAnalysisReport(
            file_path=str(target_file),
            detected_mime=mime,
            file_size_bytes=file_size,
            page_count=page_count,
            executive_summary=summary,
            key_findings=key_findings,
            extracted_data=extracted_data,
            deadlines_and_actions=deadlines,
            source_references=[f"file:///{target_file.resolve()}"],
            limitations=limitations,
            is_partial_analysis=is_partial,
        )

    def compare_files(self, file_a: Path, file_b: Path) -> Dict[str, Any]:
        """Produce structured comparative analysis between two documents."""
        rep_a = self.analyze_file(file_a)
        rep_b = self.analyze_file(file_b)

        return {
            "file_a": {
                "name": file_a.name,
                "mime": rep_a.detected_mime,
                "size_kb": round(rep_a.file_size_bytes / 1024, 1),
                "findings": rep_a.key_findings,
            },
            "file_b": {
                "name": file_b.name,
                "mime": rep_b.detected_mime,
                "size_kb": round(rep_b.file_size_bytes / 1024, 1),
                "findings": rep_b.key_findings,
            },
            "comparison_summary": f"Comparison between {file_a.name} and {file_b.name} complete. Both documents processed securely without macro/script execution.",
            "common_actions": list(set(rep_a.deadlines_and_actions).intersection(set(rep_b.deadlines_and_actions))),
        }
