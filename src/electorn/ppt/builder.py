"""PowerPoint presentation builder synthesizing editable layouts, native shapes, and typography."""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Optional, List
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from electorn.ppt.design import DesignTheme, get_theme
from electorn.ppt.schema import (
    ContentCard,
    PresentationDeckPlan,
    PresentationStyle,
    SlideContent,
    SlideLayoutType,
)


class PowerPointBuilder:
    """Renders structured PresentationDeckPlan into professional .pptx presentation decks.
    
    Adheres strictly to the requirement:
    - Genuine editable shapes, text boxes, and tables
    - Clean typography, margins, and contrast
    - High-variety slide layouts (statistics, architecture, process, timeline, tables, comparisons)
    """

    SLIDE_WIDTH = Inches(13.333)  # Standard 16:9 Widescreen
    SLIDE_HEIGHT = Inches(7.5)

    def __init__(self, style: PresentationStyle = PresentationStyle.PREMIUM_DARK) -> None:
        self.theme: DesignTheme = get_theme(style)

    def build_deck(self, plan: PresentationDeckPlan, output_path: str | Path) -> str:
        """Render all slides in the plan to disk and return absolute file path."""
        prs = Presentation()
        prs.slide_width = self.SLIDE_WIDTH
        prs.slide_height = self.SLIDE_HEIGHT

        # Blank slide layout in standard template
        blank_slide_layout = prs.slide_layouts[6]

        for slide_data in plan.slides:
            slide = prs.slides.add_slide(blank_slide_layout)
            self._render_background(slide)
            self._render_slide_content(slide, slide_data)
            self._render_speaker_notes(slide, slide_data)

        out_path = Path(output_path).resolve()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        return str(out_path)

    def _render_background(self, slide) -> None:
        """Fill slide background with theme background color."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.theme.colors.background

    def _render_speaker_notes(self, slide, data: SlideContent) -> None:
        """Attach speaker notes if provided."""
        if data.speaker_notes:
            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            tf.text = data.speaker_notes

    def _render_slide_content(self, slide, data: SlideContent) -> None:
        """Dispatch rendering based on specific layout type."""
        layout = data.layout_type

        if layout == SlideLayoutType.TITLE:
            self._render_title_slide(slide, data)
        elif layout == SlideLayoutType.SECTION_HEADER:
            self._render_section_slide(slide, data)
        elif layout == SlideLayoutType.STATISTICS:
            self._render_statistics_slide(slide, data)
        elif layout == SlideLayoutType.COMPARISON or layout == SlideLayoutType.TWO_COLUMN:
            self._render_comparison_slide(slide, data)
        elif layout == SlideLayoutType.PROCESS:
            self._render_process_slide(slide, data)
        elif layout == SlideLayoutType.ARCHITECTURE:
            self._render_architecture_slide(slide, data)
        elif layout == SlideLayoutType.TIMELINE:
            self._render_timeline_slide(slide, data)
        elif layout == SlideLayoutType.TABLE:
            self._render_table_slide(slide, data)
        elif layout == SlideLayoutType.CODE:
            self._render_code_slide(slide, data)
        elif layout == SlideLayoutType.REFERENCES:
            self._render_references_slide(slide, data)
        elif layout == SlideLayoutType.CLOSING:
            self._render_closing_slide(slide, data)
        else:
            self._render_editorial_slide(slide, data)

    # -------------------------------------------------------------
    # Shared Header & Badge Renderer
    # -------------------------------------------------------------
    def _render_header(self, slide, data: SlideContent) -> None:
        """Render badge, slide title, subtitle, and accent divider."""
        # Category Badge
        if data.category_badge:
            badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(4.0), Inches(0.35))
            btf = badge_box.text_frame
            btf.word_wrap = True
            bp = btf.paragraphs[0]
            bp.text = f"[  {data.category_badge.upper()}  ]"
            bp.font.name = self.theme.typography.title_font
            bp.font.size = self.theme.typography.caption_size
            bp.font.bold = True
            bp.font.color.rgb = self.theme.colors.accent_primary

        # Title
        title_top = Inches(0.85) if data.category_badge else Inches(0.6)
        title_box = slide.shapes.add_textbox(Inches(0.8), title_top, Inches(11.7), Inches(0.8))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.text = data.title
        tp.font.name = self.theme.typography.title_font
        tp.font.size = self.theme.typography.section_size
        tp.font.bold = True
        tp.font.color.rgb = self.theme.colors.text_primary

        # Subtitle
        if data.subtitle:
            sub_top = title_top + Inches(0.65)
            sub_box = slide.shapes.add_textbox(Inches(0.8), sub_top, Inches(11.7), Inches(0.5))
            stf = sub_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = data.subtitle
            sp.font.name = self.theme.typography.body_font
            sp.font.size = self.theme.typography.body_size
            sp.font.color.rgb = self.theme.colors.text_muted

        # Subtle Accent Bar
        bar_top = title_top + Inches(1.2) if data.subtitle else title_top + Inches(0.75)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), bar_top, Inches(1.5), Inches(0.04))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.theme.colors.accent_primary
        bar.line.fill.background()

    # -------------------------------------------------------------
    # 1. Title Slide
    # -------------------------------------------------------------
    def _render_title_slide(self, slide, data: SlideContent) -> None:
        # Decorative left accent glow / card
        accent_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.12), Inches(4.5))
        accent_card.fill.solid()
        accent_card.fill.fore_color.rgb = self.theme.colors.accent_primary
        accent_card.line.fill.background()

        # Top Badge Pill
        badge_text = (data.category_badge or "KEYNOTE TECHNICAL BRIEFING").upper()
        badge_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.5), Inches(4.5), Inches(0.4))
        badge_pill.fill.solid()
        badge_pill.fill.fore_color.rgb = self.theme.colors.surface
        badge_pill.line.color.rgb = self.theme.colors.surface_border
        bp = badge_pill.text_frame.paragraphs[0]
        bp.text = f"[  {badge_text}  ]"
        bp.font.name = self.theme.typography.title_font
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = self.theme.colors.accent_primary

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(2.0))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.text = data.title
        tp.font.name = self.theme.typography.title_font
        tp.font.size = Pt(38)
        tp.font.bold = True
        tp.font.color.rgb = self.theme.colors.text_primary

        # Subtitle
        if data.subtitle:
            sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.2), Inches(11.0), Inches(1.3))
            stf = sub_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = data.subtitle
            sp.font.name = self.theme.typography.body_font
            sp.font.size = Pt(16)
            sp.font.color.rgb = self.theme.colors.text_secondary

        # Bottom 3-Pill Executive Briefing Chips
        chips = [
            ("01 // DOMAIN", "Systems Architecture"),
            ("02 // RIGOR", "Empirical Benchmarks"),
            ("03 // VERIFICATION", "Verified Industry Literature"),
        ]
        chip_w = Inches(3.4)
        for idx, (chip_tag, chip_val) in enumerate(chips):
            chip_left = Inches(1.2) + idx * Inches(3.7)
            chip_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, chip_left, Inches(5.95), chip_w, Inches(0.55))
            chip_box.fill.solid()
            chip_box.fill.fore_color.rgb = self.theme.colors.code_bg
            chip_box.line.color.rgb = self.theme.colors.surface_border
            cp = chip_box.text_frame.paragraphs[0]
            cp.text = f"{chip_tag}: {chip_val}"
            cp.font.name = self.theme.typography.title_font
            cp.font.size = Pt(9.5)
            cp.font.bold = True
            cp.font.color.rgb = self.theme.colors.text_muted

    # -------------------------------------------------------------
    # 2. Section Header Slide
    # -------------------------------------------------------------
    def _render_section_slide(self, slide, data: SlideContent) -> None:
        self._render_editorial_slide(slide, data)

    # -------------------------------------------------------------
    # 3. Statistics & Metrics Slide
    # -------------------------------------------------------------
    def _render_statistics_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        metrics = data.metrics or [
            {"value": "95%", "label": "Benchmark Precision", "description": "Deterministic model accuracy validated across edge testbeds."},
            {"value": "3.5x", "label": "Performance Lift", "description": "Hardware-accelerated throughput compared to legacy baseline."},
            {"value": "< 100ms", "label": "Latency SLA", "description": "Real-time deadline satisfaction under peak operational load."},
        ]
        num_cards = min(4, len(metrics))
        has_bottom = bool(data.bullets or data.callout)
        card_h = Inches(3.4) if has_bottom else Inches(4.5)
        card_w = Inches((11.73 - (num_cards - 1) * 0.35) / num_cards)

        for i, m in enumerate(metrics[:num_cards]):
            if isinstance(m, dict):
                val = str(m.get("value", ""))
                lbl = str(m.get("label", ""))
                desc = str(m.get("description", ""))
            else:
                val = str(getattr(m, "value", ""))
                lbl = str(getattr(m, "label", ""))
                desc = str(getattr(m, "description", "") or "")

            left = Inches(0.8) + i * (card_w + Inches(0.35))
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.25), card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = self.theme.colors.surface
            card.line.color.rgb = self.theme.colors.accent_primary if i == 0 else self.theme.colors.surface_border

            # Number Text
            num_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.45), card_w - Inches(0.4), Inches(1.1))
            np = num_box.text_frame.paragraphs[0]
            np.text = val
            np.alignment = PP_ALIGN.CENTER
            np.font.name = self.theme.typography.title_font
            np.font.size = Pt(40)
            np.font.bold = True
            np.font.color.rgb = self.theme.colors.accent_primary if i % 2 == 0 else self.theme.colors.accent_secondary

            # Label Text
            lbl_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(3.55), card_w - Inches(0.4), Inches(0.6))
            lp = lbl_box.text_frame.paragraphs[0]
            lp.text = lbl
            lp.alignment = PP_ALIGN.CENTER
            lp.font.name = self.theme.typography.title_font
            lp.font.size = Pt(14.5)
            lp.font.bold = True
            lp.font.color.rgb = self.theme.colors.text_primary

            # Description Paragraph
            if desc:
                desc_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(4.15), card_w - Inches(0.4), card_h - Inches(2.0))
                dtf = desc_box.text_frame
                dtf.word_wrap = True
                dp = dtf.paragraphs[0]
                dp.text = desc
                dp.alignment = PP_ALIGN.CENTER
                dp.font.name = self.theme.typography.body_font
                dp.font.size = Pt(11)
                dp.font.color.rgb = self.theme.colors.text_secondary

        # Bottom Insight / Callout
        if has_bottom:
            callout_text = data.callout or (data.bullets[0] if data.bullets else "Verified telemetry metrics against empirical baselines.")
            self._render_callout_banner(slide, callout_text, "EMPIRICAL BENCHMARK TAKEAWAY")

    # -------------------------------------------------------------
    # 4. Comparison & Two-Column Slide
    # -------------------------------------------------------------
    def _render_comparison_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        col_w = Inches(5.64)
        col_h = Inches(4.5)

        bullets_l = data.left_column_bullets
        bullets_r = data.right_column_bullets
        if not bullets_l and not bullets_r:
            all_b = data.bullets or ["Legacy Baseline Paradigm", "Vulnerable Attack Vectors"]
            half = max(1, len(all_b) // 2)
            bullets_l = all_b[:half]
            bullets_r = all_b[half:] if len(all_b) > half else all_b

        # Left Column Card (Legacy / Vulnerable)
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.25), col_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = self.theme.colors.surface
        left_card.line.color.rgb = self.theme.colors.surface_border

        # Top Tag
        ltag = slide.shapes.add_textbox(Inches(1.1), Inches(2.35), col_w - Inches(0.6), Inches(0.35))
        ltp = ltag.text_frame.paragraphs[0]
        ltp.text = "[ VULNERABILITY & LEGACY BASELINE ]"
        ltp.font.name = self.theme.typography.title_font
        ltp.font.bold = True
        ltp.font.size = Pt(10)
        ltp.font.color.rgb = self.theme.colors.accent_secondary

        # Left Title
        left_title = data.left_column_title or "Legacy Architecture Limitations"
        lt_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.7), col_w - Inches(0.6), Inches(0.6))
        lp = lt_box.text_frame.paragraphs[0]
        lp.text = left_title
        lp.font.name = self.theme.typography.title_font
        lp.font.size = Pt(17)
        lp.font.bold = True
        lp.font.color.rgb = self.theme.colors.text_primary

        # Left Column Bullets
        lb_box = slide.shapes.add_textbox(Inches(1.1), Inches(3.35), col_w - Inches(0.6), col_h - Inches(1.3))
        ltf = lb_box.text_frame
        ltf.word_wrap = True
        for i, b in enumerate(bullets_l[:5]):
            bp = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
            bp.text = f"•  {b}"
            bp.font.name = self.theme.typography.body_font
            bp.font.size = Pt(12)
            bp.font.color.rgb = self.theme.colors.text_secondary
            bp.space_before = Pt(8)

        # Right Column Card (Hardened Defense / Modern Architecture)
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.89), Inches(2.25), col_w, col_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = self.theme.colors.surface
        right_card.line.color.rgb = self.theme.colors.accent_primary

        # Top Tag
        rtag = slide.shapes.add_textbox(Inches(7.19), Inches(2.35), col_w - Inches(0.6), Inches(0.35))
        rtp = rtag.text_frame.paragraphs[0]
        rtp.text = "[ HARDENED AUTONOMOUS DEFENSE ]"
        rtp.font.name = self.theme.typography.title_font
        rtp.font.bold = True
        rtp.font.size = Pt(10)
        rtp.font.color.rgb = self.theme.colors.accent_primary

        # Right Title
        right_title = data.right_column_title or "Target Autonomous Defense State"
        rt_box = slide.shapes.add_textbox(Inches(7.19), Inches(2.7), col_w - Inches(0.6), Inches(0.6))
        rp = rt_box.text_frame.paragraphs[0]
        rp.text = right_title
        rp.font.name = self.theme.typography.title_font
        rp.font.size = Pt(17)
        rp.font.bold = True
        rp.font.color.rgb = self.theme.colors.text_primary

        # Right Column Bullets
        rb_box = slide.shapes.add_textbox(Inches(7.19), Inches(3.35), col_w - Inches(0.6), col_h - Inches(1.3))
        rtf = rb_box.text_frame
        rtf.word_wrap = True
        for i, b in enumerate(bullets_r[:5]):
            bp = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
            bp.text = f"✔  {b}"
            bp.font.name = self.theme.typography.body_font
            bp.font.size = Pt(12)
            bp.font.color.rgb = self.theme.colors.text_primary
            bp.space_before = Pt(8)

    # -------------------------------------------------------------
    # 5. Process & Step Flow Slide
    # -------------------------------------------------------------
    def _render_process_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        steps = data.process_steps or [
            {"number": 1, "title": "Phase 1: Ingestion", "description": "Extract content from multi-format docs"},
            {"number": 2, "title": "Phase 2: Planning", "description": "Synthesize narrative structures"},
            {"number": 3, "title": "Phase 3: Generation", "description": "Build editable PowerPoint elements"},
            {"number": 4, "title": "Phase 4: Verification", "description": "Validate typography and bounding boxes"},
        ]
        num_steps = len(steps)
        card_w = Inches(11.7 / num_steps - 0.25)
        card_h = Inches(4.2)

        for i, s in enumerate(steps):
            if isinstance(s, dict):
                num = s.get("number", i + 1)
                title = str(s.get("title", ""))
                desc = str(s.get("description", ""))
            else:
                num = getattr(s, "number", i + 1)
                title = str(getattr(s, "title", ""))
                desc = str(getattr(s, "description", ""))

            left = Inches(0.8 + i * (11.7 / num_steps))
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.4), card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = self.theme.colors.surface
            card.line.color.rgb = self.theme.colors.surface_border

            # Number pill badge
            pill = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.2), Inches(2.6), Inches(0.55), Inches(0.55))
            pill.fill.solid()
            pill.fill.fore_color.rgb = self.theme.colors.accent_primary
            pill.line.fill.background()
            pp = pill.text_frame.paragraphs[0]
            pp.text = str(num)
            pp.alignment = PP_ALIGN.CENTER
            pp.font.bold = True
            pp.font.size = Pt(14)
            pp.font.color.rgb = RGBColor(11, 15, 25)

            # Step Title
            st_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(3.3), card_w - Inches(0.4), Inches(0.8))
            stp = st_box.text_frame.paragraphs[0]
            stp.text = title
            stp.font.name = self.theme.typography.title_font
            stp.font.size = Pt(15)
            stp.font.bold = True
            stp.font.color.rgb = self.theme.colors.text_primary

            # Step Description
            sd_box = slide.shapes.add_textbox(left + Inches(0.2), Inches(4.2), card_w - Inches(0.4), Inches(2.0))
            sdf = sd_box.text_frame
            sdf.word_wrap = True
            sdp = sdf.paragraphs[0]
            sdp.text = desc
            sdp.font.name = self.theme.typography.body_font
            sdp.font.size = Pt(12)
            sdp.font.color.rgb = self.theme.colors.text_secondary

    # -------------------------------------------------------------
    # 6. Technical Architecture Slide
    # -------------------------------------------------------------
    def _render_architecture_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        arch = data.architecture or [
            {"layer": "Layer 1: Input", "name": "Ingestion Core", "items": ["PDF", "DOCX", "Markdown"]},
            {"layer": "Layer 2: AI Core", "name": "Planning & Reasoning", "items": ["Agent Planner", "Schema Validator"]},
            {"layer": "Layer 3: Output", "name": "PowerPoint Engine", "items": ["python-pptx", "QA Verifier"]},
        ]
        layer_h = Inches(1.3)
        for i, comp in enumerate(arch):
            top = Inches(2.4 + i * 1.45)
            if isinstance(comp, dict):
                layer_name = str(comp.get("layer", ""))
                name = str(comp.get("name", ""))
                items = comp.get("items", [])
            else:
                layer_name = str(getattr(comp, "layer", ""))
                name = str(getattr(comp, "name", ""))
                items = getattr(comp, "items", [])

            # Outer Layer Box
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), layer_h)
            box.fill.solid()
            box.fill.fore_color.rgb = self.theme.colors.surface
            box.line.color.rgb = self.theme.colors.surface_border

            # Layer Header
            lh_box = slide.shapes.add_textbox(Inches(1.0), top + Inches(0.1), Inches(3.0), Inches(0.4))
            lhp = lh_box.text_frame.paragraphs[0]
            lhp.text = f"{layer_name} • {name}"
            lhp.font.bold = True
            lhp.font.size = Pt(13)
            lhp.font.color.rgb = self.theme.colors.accent_primary

            # Component sub-cards
            item_w = Inches(2.5)
            for j, item in enumerate(items[:4]):
                item_card = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(1.0 + j * 2.7),
                    top + Inches(0.55),
                    item_w,
                    Inches(0.55)
                )
                item_card.fill.solid()
                item_card.fill.fore_color.rgb = self.theme.colors.code_bg
                item_card.line.color.rgb = self.theme.colors.surface_border
                ip = item_card.text_frame.paragraphs[0]
                ip.text = item
                ip.alignment = PP_ALIGN.CENTER
                ip.font.size = Pt(12)
                ip.font.color.rgb = self.theme.colors.text_primary

    # -------------------------------------------------------------
    # 7. Timeline Slide
    # -------------------------------------------------------------
    def _render_timeline_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        items = data.timeline or [
            {"step": "Milestone 1", "title": "Discovery", "description": "System architecture and scope mapping"},
            {"step": "Milestone 2", "title": "Prototyping", "description": "Engine synthesis and testing"},
            {"step": "Milestone 3", "title": "Production", "description": "Enterprise deployment and verification"},
        ]
        # Horizontal Connecting Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.2), Inches(10.3), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = self.theme.colors.accent_primary
        line.line.fill.background()

        num_items = len(items)
        spacing = 10.3 / max(1, num_items - 1) if num_items > 1 else 5.0

        for i, item in enumerate(items):
            if isinstance(item, dict):
                step = str(item.get("step", ""))
                title = str(item.get("title", ""))
                desc = str(item.get("description", ""))
            else:
                step = str(getattr(item, "step", ""))
                title = str(getattr(item, "title", ""))
                desc = str(getattr(item, "description", ""))

            center_x = Inches(1.5 + i * spacing)

            # Node Dot
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, center_x - Inches(0.2), Inches(3.02), Inches(0.4), Inches(0.4))
            dot.fill.solid()
            dot.fill.fore_color.rgb = self.theme.colors.accent_secondary
            dot.line.color.rgb = self.theme.colors.background

            # Step label above
            top_box = slide.shapes.add_textbox(center_x - Inches(1.5), Inches(2.4), Inches(3.0), Inches(0.5))
            tp = top_box.text_frame.paragraphs[0]
            tp.text = step.upper()
            tp.alignment = PP_ALIGN.CENTER
            tp.font.bold = True
            tp.font.size = Pt(11)
            tp.font.color.rgb = self.theme.colors.accent_primary

            # Content card below
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, center_x - Inches(1.4), Inches(3.7), Inches(2.8), Inches(2.8))
            card.fill.solid()
            card.fill.fore_color.rgb = self.theme.colors.surface
            card.line.color.rgb = self.theme.colors.surface_border

            ct_box = slide.shapes.add_textbox(center_x - Inches(1.3), Inches(3.8), Inches(2.6), Inches(0.5))
            ctp = ct_box.text_frame.paragraphs[0]
            ctp.text = title
            ctp.font.bold = True
            ctp.font.size = Pt(14)
            ctp.font.color.rgb = self.theme.colors.text_primary

            cd_box = slide.shapes.add_textbox(center_x - Inches(1.3), Inches(4.3), Inches(2.6), Inches(2.0))
            cdf = cd_box.text_frame
            cdf.word_wrap = True
            cdp = cdf.paragraphs[0]
            cdp.text = desc
            cdp.font.size = Pt(11)
            cdp.font.color.rgb = self.theme.colors.text_secondary

    # -------------------------------------------------------------
    # 8. Table Slide
    # -------------------------------------------------------------
    def _render_table_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        table_data = data.table

        headers = table_data.headers if table_data else ["Criteria", "Traditional Method", "Electorn Engine"]
        rows = table_data.rows if table_data else [
            ["Generation Speed", "Hours manual", "Seconds autonomous"],
            ["Verification", "Manual review", "Automated QA inspector"],
            ["Editable Shapes", "No (Screenshots)", "Yes (Native Vector)"],
        ]

        num_rows = len(rows) + 1
        num_cols = len(headers)
        table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(0.8), Inches(2.4), Inches(11.7), Inches(4.4))
        tbl = table_shape.table

        # Format Header Row
        for col_idx, h in enumerate(headers):
            cell = tbl.cell(0, col_idx)
            cell.text = str(h)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.theme.colors.surface
            p = cell.text_frame.paragraphs[0]
            p.font.name = self.theme.typography.title_font
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = self.theme.colors.accent_primary

        # Format Data Rows
        for row_idx, r in enumerate(rows):
            for col_idx, val in enumerate(r[:num_cols]):
                cell = tbl.cell(row_idx + 1, col_idx)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.theme.colors.code_bg
                p = cell.text_frame.paragraphs[0]
                p.font.name = self.theme.typography.body_font
                p.font.size = Pt(12)
                p.font.color.rgb = self.theme.colors.text_primary

    # -------------------------------------------------------------
    # 9. Code Slide
    # -------------------------------------------------------------
    def _render_code_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        code_text = data.code_snippet or 'print("Electorn AI autonomous execution active")'

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(11.7), Inches(4.4))
        card.fill.solid()
        card.fill.fore_color.rgb = self.theme.colors.code_bg
        card.line.color.rgb = self.theme.colors.surface_border

        code_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(11.1), Inches(4.0))
        ctf = code_box.text_frame
        ctf.word_wrap = True
        lines = code_text.split("\n")
        for i, line in enumerate(lines):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            p.text = line
            p.font.name = self.theme.typography.code_font
            p.font.size = Pt(12)
            p.font.color.rgb = self.theme.colors.text_primary

    # -------------------------------------------------------------
    # 10. References Slide
    # -------------------------------------------------------------
    def _render_references_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        citations = data.citations or [
            "Official Electorn AI Documentation & Technical Blueprint v0.1.0",
            "IEEE Autonomous Agents and Multi-Tier Orchestration Standards",
            "W3C Architecture Recommendations for Automation Runtimes",
        ]

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.4), Inches(11.7), Inches(4.4))
        card.fill.solid()
        card.fill.fore_color.rgb = self.theme.colors.surface
        card.line.color.rgb = self.theme.colors.surface_border

        ref_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(11.0), Inches(3.8))
        rtf = ref_box.text_frame
        rtf.word_wrap = True
        for i, cite in enumerate(citations):
            p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
            p.text = f"[{i+1}]  {cite}"
            p.font.name = self.theme.typography.body_font
            p.font.size = Pt(14)
            p.font.color.rgb = self.theme.colors.text_secondary
            p.space_before = Pt(14)

    # -------------------------------------------------------------
    # 11. Closing Slide (Gamma Executive Synthesis)
    # -------------------------------------------------------------
    def _render_closing_slide(self, slide, data: SlideContent) -> None:
        self._render_header(slide, data)
        
        cards = data.cards
        if not cards and data.bullets:
            cards = self._parse_bullets_into_cards(data.bullets)
            
        if not cards:
            cards = [
                ContentCard(
                    tag="01 // ARCHITECTURE",
                    title="Defense Posture Integration",
                    body="Formally verified defense combining hardware root of trust, cryptographic in-vehicle bus signing, and runtime telemetry monitoring.",
                ),
                ContentCard(
                    tag="02 // STANDARDS",
                    title="Regulatory Alignment",
                    body="Full lifecycle alignment with ISO/SAE 21434 and UNECE WP.29 regulatory audit frameworks for production fleets.",
                ),
                ContentCard(
                    tag="03 // DEPLOYMENT",
                    title="Production Action Plan",
                    body="Execute threat analysis and risk assessment (TARA) followed by phased deployment and continuous over-the-air observability.",
                ),
            ]
            
        self._render_three_card_grid(slide, cards[:3], Inches(4.5))

    # -------------------------------------------------------------
    # 12. Editorial / General Slide (Gamma Multi-Card Layouts)
    # -------------------------------------------------------------
    def _render_editorial_slide(self, slide, data: SlideContent) -> None:
        """Render Gamma-style modern structured cards layout with callouts."""
        self._render_header(slide, data)

        cards = data.cards
        if not cards and data.bullets:
            cards = self._parse_bullets_into_cards(data.bullets)

        if not cards:
            cards = [
                ContentCard(
                    tag="01 // ARCHITECTURE",
                    title="Core Principles & Mechanism",
                    body="Comprehensive evaluation of system foundations, operational constraints, and telemetry parameters.",
                ),
                ContentCard(
                    tag="02 // EXECUTION",
                    title="Empirical Validation",
                    body="Deterministic verification confirming baseline benchmarks, threat resistance, and runtime stability.",
                ),
            ]

        has_callout = bool(data.callout)
        callout_h = Inches(1.15) if has_callout else Inches(0)
        available_h = Inches(4.5) - callout_h

        num_cards = len(cards)
        if num_cards <= 2:
            self._render_two_card_grid(slide, cards, available_h)
        elif num_cards == 3:
            self._render_three_card_grid(slide, cards, available_h)
        else:
            self._render_four_card_grid(slide, cards[:4], available_h)

        if has_callout:
            self._render_callout_banner(slide, data.callout, data.callout_label)

    # -------------------------------------------------------------
    # Gamma Card Grid Helpers
    # -------------------------------------------------------------
    def _parse_bullets_into_cards(self, bullets: list[str]) -> list[ContentCard]:
        """Convert standard bullet strings into rich Gamma structured cards."""
        cards = []
        for i, b in enumerate(bullets):
            b_clean = b.strip().lstrip("•-*✔ ").strip()
            if not b_clean:
                continue
            # Check for colon or dash delimiter: "Title: Body"
            parts = re.split(r"\s*[:—–]\s*", b_clean, maxsplit=1)
            if len(parts) == 2 and len(parts[0]) < 65:
                title = parts[0].strip()
                body = parts[1].strip()
            else:
                words = b_clean.split()
                if len(words) > 5:
                    title = " ".join(words[:4]).strip()
                    body = b_clean
                else:
                    title = b_clean
                    body = b_clean

            tag = f"0{i+1} // PILLAR"
            cards.append(ContentCard(tag=tag, title=title, body=body))
        return cards

    def _render_two_card_grid(self, slide, cards: list[ContentCard], card_h: Inches) -> None:
        """Render 2 wide side-by-side Gamma cards."""
        col_w = Inches(5.64)
        top = Inches(2.25)
        for i, card in enumerate(cards[:2]):
            left = Inches(0.8) + i * Inches(6.09)
            # Outer Card
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_w, card_h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.colors.surface
            shape.line.color.rgb = self.theme.colors.accent_primary if i == 1 else self.theme.colors.surface_border

            # Tag
            tag_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), col_w - Inches(0.6), Inches(0.35))
            tag_p = tag_box.text_frame.paragraphs[0]
            tag_p.text = f"[  {(card.tag or f'0{i+1} // ANALYSIS').upper()}  ]"
            tag_p.font.name = self.theme.typography.title_font
            tag_p.font.size = Pt(10)
            tag_p.font.bold = True
            tag_p.font.color.rgb = self.theme.colors.accent_primary if i == 0 else self.theme.colors.accent_secondary

            # Title
            t_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.55), col_w - Inches(0.6), Inches(0.7))
            t_box.text_frame.word_wrap = True
            tp = t_box.text_frame.paragraphs[0]
            tp.text = card.title
            tp.font.name = self.theme.typography.title_font
            tp.font.size = Pt(17)
            tp.font.bold = True
            tp.font.color.rgb = self.theme.colors.text_primary

            # Accent line under title
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.3), top + Inches(1.25), Inches(1.2), Inches(0.02))
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.theme.colors.accent_primary if i == 1 else self.theme.colors.accent_secondary
            bar.line.fill.background()

            # Body
            b_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(1.35), col_w - Inches(0.6), card_h - Inches(1.5))
            btf = b_box.text_frame
            btf.word_wrap = True
            bp = btf.paragraphs[0]
            bp.text = card.body
            bp.font.name = self.theme.typography.body_font
            bp.font.size = Pt(12.5)
            bp.font.color.rgb = self.theme.colors.text_secondary
            bp.space_before = Pt(4)

            # Secondary bullets if provided
            if card.bullet_points:
                for bullet_txt in card.bullet_points[:3]:
                    sub_p = btf.add_paragraph()
                    sub_p.text = f"•  {bullet_txt}"
                    sub_p.font.name = self.theme.typography.body_font
                    sub_p.font.size = Pt(11)
                    sub_p.font.color.rgb = self.theme.colors.text_muted
                    sub_p.space_before = Pt(6)

    def _render_three_card_grid(self, slide, cards: list[ContentCard], card_h: Inches) -> None:
        """Render 3 horizontal Gamma cards side-by-side."""
        card_w = Inches(3.67)
        top = Inches(2.25)
        for i, card in enumerate(cards[:3]):
            left = Inches(0.8) + i * Inches(4.03)
            # Outer Card
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.colors.surface
            shape.line.color.rgb = self.theme.colors.accent_primary if i == 1 else self.theme.colors.surface_border

            # Tag
            tag_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), card_w - Inches(0.5), Inches(0.35))
            tag_p = tag_box.text_frame.paragraphs[0]
            tag_p.text = f"[  {(card.tag or f'0{i+1} // FOCUS').upper()}  ]"
            tag_p.font.name = self.theme.typography.title_font
            tag_p.font.size = Pt(9.5)
            tag_p.font.bold = True
            tag_p.font.color.rgb = self.theme.colors.accent_primary

            # Title
            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.55), card_w - Inches(0.5), Inches(0.75))
            t_box.text_frame.word_wrap = True
            tp = t_box.text_frame.paragraphs[0]
            tp.text = card.title
            tp.font.name = self.theme.typography.title_font
            tp.font.size = Pt(15)
            tp.font.bold = True
            tp.font.color.rgb = self.theme.colors.text_primary

            # Accent line under title
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.25), top + Inches(1.3), Inches(0.8), Inches(0.02))
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.theme.colors.accent_primary
            bar.line.fill.background()

            # Body
            b_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(1.4), card_w - Inches(0.5), card_h - Inches(1.5))
            btf = b_box.text_frame
            btf.word_wrap = True
            bp = btf.paragraphs[0]
            bp.text = card.body
            bp.font.name = self.theme.typography.body_font
            bp.font.size = Pt(11.5)
            bp.font.color.rgb = self.theme.colors.text_secondary

    def _render_four_card_grid(self, slide, cards: list[ContentCard], card_h: Inches) -> None:
        """Render 4 Gamma cards in a balanced 2x2 Bento-box grid."""
        card_w = Inches(5.64)
        single_h = (card_h - Inches(0.25)) / 2
        for i, card in enumerate(cards[:4]):
            row = i // 2
            col = i % 2
            left = Inches(0.8) + col * Inches(6.09)
            top = Inches(2.25) + row * (single_h + Inches(0.25))

            # Outer Card
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, single_h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.colors.surface
            shape.line.color.rgb = self.theme.colors.surface_border

            # Tag
            tag_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.12), card_w - Inches(0.5), Inches(0.3))
            tag_p = tag_box.text_frame.paragraphs[0]
            tag_p.text = f"[  {(card.tag or f'0{i+1} // PILLAR').upper()}  ]"
            tag_p.font.name = self.theme.typography.title_font
            tag_p.font.size = Pt(9)
            tag_p.font.bold = True
            tag_p.font.color.rgb = self.theme.colors.accent_primary

            # Title
            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.4), card_w - Inches(0.5), Inches(0.45))
            t_box.text_frame.word_wrap = True
            tp = t_box.text_frame.paragraphs[0]
            tp.text = card.title
            tp.font.name = self.theme.typography.title_font
            tp.font.size = Pt(13.5)
            tp.font.bold = True
            tp.font.color.rgb = self.theme.colors.text_primary

            # Body
            b_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.85), card_w - Inches(0.5), single_h - Inches(0.9))
            btf = b_box.text_frame
            btf.word_wrap = True
            bp = btf.paragraphs[0]
            bp.text = card.body
            bp.font.name = self.theme.typography.body_font
            bp.font.size = Pt(10.5)
            bp.font.color.rgb = self.theme.colors.text_secondary

    def _render_callout_banner(self, slide, callout_text: str, callout_label: Optional[str] = None) -> None:
        """Render a sleek, high-contrast callout banner across the bottom."""
        left = Inches(0.8)
        top = Inches(6.0)
        width = Inches(11.73)
        height = Inches(0.95)

        banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        banner.fill.solid()
        banner.fill.fore_color.rgb = self.theme.colors.code_bg
        banner.line.color.rgb = self.theme.colors.accent_primary

        tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.12), width - Inches(0.6), height - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        label = (callout_label or "STRATEGIC IMPERATIVE").upper()
        p.text = f"⚡ {label}: "
        p.font.name = self.theme.typography.title_font
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = self.theme.colors.accent_primary

        run = p.add_run()
        run.text = callout_text
        run.font.name = self.theme.typography.body_font
        run.font.size = Pt(11)
        run.font.bold = False
        run.font.color.rgb = self.theme.colors.text_primary
