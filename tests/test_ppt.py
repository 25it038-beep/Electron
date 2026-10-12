import pytest
from pathlib import Path
from typer.testing import CliRunner
import pptx

from electorn.ppt.schema import (
    PresentationDeckPlan,
    SlideContent,
    SlideLayoutType,
    PresentationStyle,
    KeyMetric,
    TimelineItem,
    ProcessStep,
    ArchComponent,
    TableData,
)
from electorn.ppt.design import THEMES, get_theme
from electorn.ppt.document_ingestion import DocumentIngestionEngine
from electorn.ppt.planner import PresentationPlanner
from electorn.ppt.builder import PowerPointBuilder
from electorn.ppt.validator import PresentationQAValidator
from electorn.ppt.engine import PresentationEngine
from electorn.cli import app

runner = CliRunner()


def test_theme_retrieval():
    """Verify all 7 presentation styles exist and have valid palettes and typography."""
    for style in PresentationStyle:
        theme = get_theme(style)
        assert theme is not None
        assert theme.colors.background is not None
        assert theme.colors.surface is not None
        assert theme.colors.accent_primary is not None
        assert theme.typography.title_size.pt > 0
        assert theme.typography.body_size.pt > 0


def test_document_ingestion_markdown(tmp_path: Path):
    """Verify markdown ingestion extracts headers, text, and bullet points."""
    sample_md = tmp_path / "notes.md"
    sample_md.write_text(
        "# Executive Summary\n"
        "AI is revolutionizing healthcare diagnostics.\n\n"
        "## Key Findings\n"
        "- 98% accuracy in image classification\n"
        "- Reduced triage latency by 45%\n\n"
        "## Architecture\n"
        "Input layer feeds into transformer encoder.\n",
        encoding="utf-8"
    )

    doc_data = DocumentIngestionEngine.ingest_file(sample_md)

    assert doc_data.source_file == "notes.md"
    assert "Markdown" in doc_data.file_type
    assert len(doc_data.sections) >= 2
    assert any("Key Findings" in s.title or "Executive Summary" in s.title for s in doc_data.sections)
    assert doc_data.total_characters > 50


def test_planner_narrative_generation():
    """Verify the presentation planner produces a structured deck plan with diverse layouts."""
    planner = PresentationPlanner()
    plan = planner.plan_presentation(
        topic="Cloud Architecture Best Practices",
        target_slide_count=8,
        style=PresentationStyle.CORPORATE_PITCH,
        audience="Senior Engineering Leaders"
    )

    assert isinstance(plan, PresentationDeckPlan)
    assert len(plan.slides) == 8
    assert plan.slides[0].layout_type == SlideLayoutType.TITLE
    assert plan.slides[-1].layout_type == SlideLayoutType.CLOSING

    layouts = {s.layout_type for s in plan.slides}
    assert len(layouts) >= 4, "Plan should utilize diverse layout types"


def test_powerpoint_builder_and_validator(tmp_path: Path):
    """Verify builder produces a valid 16:9 PPTX deck and QA validator passes."""
    deck_plan = PresentationDeckPlan(
        topic="Microservices Architecture Review",
        audience="Senior Engineering Leaders",
        target_slide_count=8,
        style=PresentationStyle.PREMIUM_DARK,
        slides=[
            SlideContent(
                slide_number=1,
                layout_type=SlideLayoutType.TITLE,
                title="Microservices Architecture Review",
                subtitle="Scalability, Resilience, and Observability",
                speaker_notes="Welcome leadership and technical teams."
            ),
            SlideContent(
                slide_number=2,
                layout_type=SlideLayoutType.STATISTICS,
                title="System Performance Benchmarks",
                subtitle="Telemetry insights over 12 months",
                metrics=[
                    KeyMetric(value="150K RPS", label="Throughput", description="Peak load"),
                    KeyMetric(value="4.2 ms", label="P99 Latency", description="Global average"),
                    KeyMetric(value="99.995%", label="Uptime", description="High availability")
                ],
                speaker_notes="Key performance gains achieved after decoupling the monolith."
            ),
            SlideContent(
                slide_number=3,
                layout_type=SlideLayoutType.PROCESS,
                title="Deployment Pipeline",
                subtitle="Automated CD lifecycle",
                process_steps=[
                    ProcessStep(number=1, title="Build", description="Compile and containerize"),
                    ProcessStep(number=2, title="Verify", description="Contract and integration tests"),
                    ProcessStep(number=3, title="Deploy", description="Canary rollout to clusters")
                ],
                speaker_notes="Explain canary rollback triggers."
            ),
            SlideContent(
                slide_number=4,
                layout_type=SlideLayoutType.ARCHITECTURE,
                title="Distributed Mesh Topology",
                architecture=[
                    ArchComponent(layer="Ingress", name="Envoy Gateway", items=["TLS termination & routing"]),
                    ArchComponent(layer="Mesh", name="Istio Control Plane", items=["mTLS and policy"]),
                    ArchComponent(layer="Storage", name="Distributed Raft DB", items=["Multi-region state"])
                ],
                speaker_notes="Walk through layer security constraints."
            ),
            SlideContent(
                slide_number=5,
                layout_type=SlideLayoutType.TWO_COLUMN,
                title="Monolith vs Microservices",
                left_column_title="Monolithic Pattern",
                left_column_bullets=["Single failure domain", "Coupled releases", "Complex builds"],
                right_column_title="Microservices Mesh",
                right_column_bullets=["Independent deploys", "Fault isolation", "Polyglot services"],
                speaker_notes="Highlight organizational boundaries aligning with services."
            ),
            SlideContent(
                slide_number=6,
                layout_type=SlideLayoutType.CODE,
                title="Health Probe Implementation",
                code_snippet="apiVersion: v1\nkind: Pod\nspec:\n  containers:\n  - name: api\n    livenessProbe:\n      httpGet:\n        path: /healthz\n        port: 8080",
                code_language="yaml",
                bullets=["Configured automated liveness probes", "Configured Readiness gates"],
                speaker_notes="Standard Kubernetes health checking contract."
            ),
            SlideContent(
                slide_number=7,
                layout_type=SlideLayoutType.TABLE,
                title="Technology Evaluation",
                table=TableData(
                    headers=["Service", "Protocol", "Latency SLA"],
                    rows=[
                        ["Auth Service", "gRPC", "< 2ms"],
                        ["Billing API", "HTTPS / JSON", "< 15ms"],
                        ["Notification", "Kafka Stream", "Async"]
                    ]
                ),
                speaker_notes="Review SLA compliance."
            ),
            SlideContent(
                slide_number=8,
                layout_type=SlideLayoutType.CLOSING,
                title="Execution Roadmap & Next Steps",
                bullets=["Finalize cluster migration", "Engage security audit", "Open Q&A"],
                speaker_notes="Open the floor for questions."
            )
        ]
    )

    output_pptx = tmp_path / "microservices.pptx"
    builder = PowerPointBuilder(style=PresentationStyle.PREMIUM_DARK)
    saved_path = builder.build_deck(deck_plan, output_pptx)

    assert output_pptx.exists()
    assert output_pptx.stat().st_size > 5000

    # Verify widescreen 16:9
    prs = pptx.Presentation(str(output_pptx))
    assert len(prs.slides) == 8
    assert abs(prs.slide_width.inches - 13.333) < 0.1
    assert abs(prs.slide_height.inches - 7.5) < 0.1

    # Verify speaker notes are populated
    first_slide_notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Welcome leadership" in first_slide_notes

    # Verify QA validator
    qa_report = PresentationQAValidator.validate_file(saved_path, expected_slide_count=8)
    assert qa_report.is_valid is True
    assert qa_report.total_slides == 8
    assert len(qa_report.issues) == 0
    assert "slides" in qa_report.manifest


def test_presentation_engine_preview_and_inspect(tmp_path: Path):
    """Verify presentation engine inspection and preview."""
    output_pptx = tmp_path / "demo.pptx"
    engine = PresentationEngine()
    
    # Generate 5-slide deck
    saved_file, qa_report = engine.generate(
        topic="Autonomous AI Agents",
        slides_count=5,
        style="cinematic",
        output_file=output_pptx
    )

    assert Path(saved_file).exists()
    assert qa_report.is_valid is True

    # Test preview (should not raise error)
    engine.preview(saved_file)

    # Test inspect
    inspection = engine.inspect(saved_file)
    assert inspection.total_slides == 5
    assert inspection.manifest["aspect_ratio"] == "16:9 Widescreen"
    assert len(inspection.manifest["slides"]) == 5


def test_presentation_engine_revision(tmp_path: Path):
    """Verify revising a slide modifies the deck and updates notes."""
    output_pptx = tmp_path / "original.pptx"
    engine = PresentationEngine()

    saved_file, qa_report = engine.generate(
        topic="Distributed Systems",
        slides_count=4,
        style="minimalist-light",
        output_file=output_pptx
    )

    revised_file = engine.revise(
        file_path=saved_file,
        slide_number=2,
        instruction="Focus specifically on Paxos and Raft consensus algorithms with 99.999% reliability"
    )

    assert Path(revised_file).exists()
    # Verify slide 2 got updated notes instruction
    prs = pptx.Presentation(str(revised_file))
    assert len(prs.slides) == 4
    slide_2_notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "Paxos and Raft" in slide_2_notes


def test_cli_ppt_generation(tmp_path: Path):
    """Verify CLI command electorn ppt executes end-to-end."""
    output_pptx = tmp_path / "cli_test.pptx"
    result = runner.invoke(app, [
        "ppt",
        "Modern Machine Learning in Healthcare",
        "--slides", "4",
        "--style", "academic",
        "--output", str(output_pptx)
    ])
    assert result.exit_code == 0
    assert output_pptx.exists()
    assert "PowerPoint Deck Generated" in result.stdout
